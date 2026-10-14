# -*- coding: utf-8 -*-
"""관내여비 지급명세서 — 규칙·파일 해석·엑셀 만들기 (화면과 무관한 순수 코드)

 · 일비 규칙: 같은 날 1건이면 4시간 이상 전액(1일 한도), 1시간 이상 절반, 그 미만 0원.
              2건 이상이면 건당 1시간 이상인 건마다 절반씩(합계는 1일 한도를 넘지 않음).
 · 차량 감액: 차량배차 파일에 '관내 목적지' 배차 기록이 있는 날만 감액.
 · 월 한도: 금액이 큰 날부터 채우고 남는 한도를 반액인 날로 메워 한도를 최대한 소진.
"""
import csv
import datetime as dt
import re
from dataclasses import dataclass, field
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border

from ...core import excel as X

ABSOLUTE_MONTHLY_CAP = 300000

BANK_CODES = [
    "001:한국은행", "002:산업은행", "003:기업은행", "004:국민은행", "007:수협은행", "008:수출입은행",
    "011:농협은행", "012:지역 농축협", "020:우리은행", "023:SC제일은행", "027:한국씨티은행", "031:대구은행",
    "032:부산은행", "034:광주은행", "035:제주은행", "037:전북은행", "039:경남은행", "045:새마을금고중앙회",
    "048:신협", "050:상호저축은행", "071:우체국", "081:KEB하나은행", "088:신한은행", "089:케이뱅크",
    "090:카카오뱅크", "092:토스뱅크", "093:한국주택금융공사", "094:서울보증보험",
]

NEGATIVE = ["취소", "반려", "미승인", "미배차", "불가", "미운행", "운행전", "대기중", "철회"]
PLACEHOLDER = ["없음", "해당없음", "-", ".", ""]


@dataclass
class Rules:
    monthly_cap: int = 140000
    daily_cap: int = 20000
    full_hours: float = 4
    split_hours: float = 1
    vehicle_deduction: int = 10000
    outside_keywords: list = field(default_factory=lambda: ["관외"])

    def normalized(self):
        r = Rules(**self.__dict__)
        r.monthly_cap = min(int(r.monthly_cap or ABSOLUTE_MONTHLY_CAP), ABSOLUTE_MONTHLY_CAP)
        r.daily_cap = int(r.daily_cap or 20000)
        r.full_hours = float(r.full_hours or 4)
        r.split_hours = float(r.split_hours or 1)
        r.vehicle_deduction = int(r.vehicle_deduction or 0)
        r.outside_keywords = [k for k in (r.outside_keywords or []) if k] or ["관외"]
        return r


@dataclass
class Employee:
    id: str
    name: str
    dept: str = ""
    purpose_counts: dict = field(default_factory=dict)
    records: list = field(default_factory=list)       # {date, day, hours}
    all_count: int = 0

    def common_purpose(self) -> str:
        best, best_n = "", -1
        for p, n in self.purpose_counts.items():
            if n > best_n:
                best, best_n = p, n
        return best


@dataclass
class Result:
    total: int = 0
    included_days: list = field(default_factory=list)
    all_days: list = field(default_factory=list)
    capped: bool = False
    vehicle_days: list = field(default_factory=list)
    vehicle_total: int = 0


@dataclass
class Override:
    bank: str = ""
    account: str = ""
    holder: str = ""
    purpose: str = ""


class ParseError(Exception):
    pass


def cell_text(v) -> str:
    """JS 의 raw:false(화면 표시 문자열)와 비슷하게 셀을 글자로"""
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, dt.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _label(v) -> str:
    return cell_text(v).strip().replace("\r", "").replace("\n", "")


def sheet_rows(sheets: dict, which: str = "first") -> list:
    names = list(sheets)
    name = names[0] if which == "first" else names[-1]
    return [[cell_text(c) for c in r] for r in sheets[name]]


# ───────────────────── 원본(출장 집계표) 해석 ─────────────────────
def _parse_name(cell: str):
    m = re.match(r"^(.+?)\(([^)]+)\)\s*$", (cell or "").strip())
    return (m.group(1).strip(), m.group(2).strip()) if m else None


def _parse_dt(s: str):
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})\s+(\d{2}):(\d{2})", s)
    return dt.datetime(*map(int, m.groups())) if m else None


def _parse_period(cell: str):
    clean = (cell or "").replace("\r", " ").replace("\n", " ").strip()
    if "~" not in clean:
        return None
    parts = clean.split("~")
    if len(parts) != 2:
        return None
    a, b = _parse_dt(parts[0].strip()), _parse_dt(parts[1].strip())
    return (a, b) if a and b else None


def parse_employees(rows: list) -> dict:
    """반환: {id: Employee} (등장 순서 유지)"""
    cols = None
    for ri, row in enumerate(rows):
        labels = [_label(v) for v in row]
        if "성명" in labels and "출장기간" in labels:
            def idx(n):
                return labels.index(n) if n in labels else -1
            cols = dict(name=idx("성명"), period=idx("출장기간"), dept=idx("소속/직급"), purpose=idx("출장목적"))
            break
    if not cols:
        raise ParseError('엑셀에서 "성명"/"출장기간" 헤더를 찾을 수 없습니다. 파일 형식을 확인해주세요.')
    out: dict = {}
    for row in rows:
        def get(i):
            return row[i] if 0 <= i < len(row) else ""
        pn = _parse_name(get(cols["name"]))
        if not pn:
            continue
        per = _parse_period(get(cols["period"]))
        if not per:
            continue
        dur = (per[1] - per[0]).total_seconds() / 3600
        if dur < 0:
            continue
        d = per[0]
        date_key = f"{d.year}-{d.month:02d}-{d.day:02d}"
        name, pid = pn
        if pid not in out:
            dept = get(cols["dept"]) if cols["dept"] >= 0 else ""
            out[pid] = Employee(pid, name, re.sub(r"[\r\n]+", " ", dept.replace("_x000D_", " ")).strip())
        e = out[pid]
        e.all_count += 1
        e.records.append({"date": date_key, "day": d.day, "hours": dur})
        if cols["purpose"] >= 0:
            p = re.sub(r"[\r\n]+", " ", get(cols["purpose"]).replace("_x000D_", " ")).strip()
            if p:
                e.purpose_counts[p] = e.purpose_counts.get(p, 0) + 1
    return out


# ───────────────────── 선택 파일: 지급대상자 목록 / 차량배차 ─────────────────────
def parse_order_file(rows: list):
    """('성명','연번' 헤더) → (names, {이름:{bank,account}}) 또는 None"""
    hdr = nidx = bidx = aidx = -1
    for r, row in enumerate(rows):
        lab = [_label(v) for v in row]
        if "성명" in lab and "연번" in lab:
            hdr, nidx = r, lab.index("성명")
            bidx = lab.index("은행명") if "은행명" in lab else -1
            aidx = lab.index("계좌번호") if "계좌번호" in lab else -1
            break
    if hdr < 0:
        return None
    names, bank_map = [], {}
    for row in rows[hdr + 1:]:
        if nidx >= len(row):
            continue
        name = row[nidx].strip()
        if not name or name == "계":
            continue
        names.append(name)
        bank = row[bidx].strip() if 0 <= bidx < len(row) else ""
        acc = row[aidx].strip() if 0 <= aidx < len(row) else ""
        if bank or acc:
            bank_map[name] = {"bank": bank, "account": acc}
    return (names, bank_map) if names else None


def read_csv_rows(path: str) -> list:
    raw = open(path, "rb").read()
    try:
        text = raw.decode("cp949")
    except UnicodeDecodeError:
        text = raw.decode("utf-8", errors="replace")
    if text.count("�") > 3:
        text = raw.decode("utf-8", errors="replace")
    return [r for r in csv.reader(text.splitlines())]


def _split_people(cell: str) -> list:
    raw = (cell or "").replace("\r", " ").replace("\n", " ").strip()
    if not raw:
        return []
    out = []
    for s in re.split(r"[-,·/、;]+", raw):
        t = s.strip()
        pn = _parse_name(t)
        n = pn[0] if pn else t
        if n and n not in PLACEHOLDER:
            out.append(n)
    return out


def _vehicle_dates(cell: str) -> list:
    raw = (cell or "").replace("\r", " ").replace("\n", " ").strip()
    if not raw:
        return []
    parts = raw.split("~") if "~" in raw else [raw]

    def ymd(s):
        m = re.search(r"(\d{4})[.\-/](\d{1,2})[.\-/](\d{1,2})", s)
        return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3))) if m else None
    try:
        start = ymd(parts[0])
        if not start:
            return []
        end = (ymd(parts[1]) if len(parts) > 1 else None) or start
    except ValueError:
        return []
    if end < start:
        return [start.strftime("%Y-%m-%d")]
    keys, cur, guard = [], start, 0
    while cur <= end and guard < 62:
        keys.append(cur.strftime("%Y-%m-%d"))
        cur += dt.timedelta(days=1)
        guard += 1
    return keys


def parse_vehicle_file(rows: list):
    """배차대장 → {'이름|날짜': [목적지, …]} (없으면 None)"""
    DATE = ["사용일시", "배차일자", "운행일자", "사용일자", "출장일자", "배차일", "운행일", "사용일", "날짜", "일자", "배차기간", "운행기간"]
    NAME = ["사용자", "성명", "신청자", "신청인", "운전자", "이름"]
    hdr = -1
    di = ni = dr = cp = st = asg = dest = -1
    for r, row in enumerate(rows):
        lab = [_label(v) for v in row]
        d = next((lab.index(x) for x in DATE if x in lab), -1)
        if d < 0:
            continue
        n = next((lab.index(x) for x in NAME if x in lab), -1)
        if n < 0:
            continue
        hdr, di, ni = r, d, n
        dr = lab.index("운전자") if "운전자" in lab else -1
        cp = lab.index("동승자") if "동승자" in lab else -1
        st = lab.index("운행상황") if "운행상황" in lab else -1
        asg = lab.index("배차여부") if "배차여부" in lab else -1
        dest = lab.index("목적지") if "목적지" in lab else -1
        break
    if hdr < 0:
        return None
    usage: dict = {}
    for row in rows[hdr + 1:]:
        def g(i):
            return row[i] if 0 <= i < len(row) else ""
        if not g(di).strip():
            continue
        s1, s2 = g(asg).strip(), g(st).strip()
        if any(k in s1 for k in NEGATIVE) or any(k in s2 for k in NEGATIVE):
            continue
        names = set(_split_people(g(ni)))
        if dr >= 0:
            names |= set(_split_people(g(dr)))
        if cp >= 0:
            names |= set(_split_people(g(cp)))
        if not names:
            continue
        dest_text = g(dest).strip() if dest >= 0 else ""
        for dk in _vehicle_dates(g(di)):
            for n in names:
                usage.setdefault(f"{n}|{dk}", []).append(dest_text)
    return usage or None


# ───────────────────── 계산 ─────────────────────
def compute(emp: Employee, rules: Rules, vehicle: dict | None) -> Result:
    rules = rules.normalized()
    by_date: dict = {}
    for r in emp.records:
        by_date.setdefault(r["date"], []).append(r)
    dates = sorted(by_date)
    day_all = []
    for d in dates:
        recs = by_date[d]
        if len(recs) == 1:
            h = recs[0]["hours"]
            fee = rules.daily_cap if h >= rules.full_hours else (rules.daily_cap / 2 if h >= rules.split_hours else 0)
        else:
            fee = 0
            for r in recs:
                if r["hours"] >= rules.split_hours:
                    fee += rules.daily_cap / 2
            fee = min(fee, rules.daily_cap)
        ded, ded_amt = False, 0
        if vehicle is not None and rules.vehicle_deduction > 0 and fee > 0:
            dests = vehicle.get(f"{emp.name}|{d}")
            if dests and any(not (t and any(k in t for k in rules.outside_keywords)) for t in dests):
                before = fee
                fee = max(0, fee - rules.vehicle_deduction)
                ded, ded_amt = True, before - fee
        day_all.append({"date": d, "day": recs[0]["day"], "fee": fee, "ded": ded, "ded_amt": ded_amt})
    day_fees = [x for x in day_all if x["fee"] > 0]
    total, included, skipped = 0, set(), False
    for item in sorted(day_fees, key=lambda x: (-x["fee"], x["date"])):
        if total + item["fee"] <= rules.monthly_cap:
            total += item["fee"]
            included.add(item["date"])
        else:
            skipped = True
    res = Result()
    res.total = int(total)
    res.included_days = [x["day"] for x in day_fees if x["date"] in included]
    res.all_days = [by_date[d][0]["day"] for d in dates]
    res.capped = skipped
    res.vehicle_days = [x["day"] for x in day_all if x["ded"]]
    res.vehicle_total = int(sum(x["ded_amt"] for x in day_all if x["ded"]))
    return res


def bank_to_code(text: str) -> str:
    t = (text or "").strip()
    if not t:
        return ""
    if re.match(r"^\d{3}:", t):
        return t

    def norm(s):
        return re.sub(r"은행|중앙회|\s", "", s)
    target = norm(t)
    if not target:
        return ""
    best = None
    for entry in BANK_CODES:
        name = norm(entry.split(":")[1])
        if name == target:
            return entry
        if best is None and (target in name or name in target):
            best = entry
    return best or t


# ───────────────────── 엑셀 만들기 ─────────────────────
NUM_FMT = '_-* #,##0_-;-* #,##0_-;_-* "-"_-;_-@_-'


def build_ledger(title: str, ap_rank: str, ap_name: str, rows: list) -> Workbook:
    """rows: [(성명, 금액, 은행, 계좌, 출장일 목록(str), 목적)] — 지급명세서(총괄 서식)"""
    from openpyxl.styles import Side
    thin, medium = Side(style="thin"), Side(style="medium")
    grand = sum(r[1] for r in rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "지급명세서"
    ws.sheet_view.showGridLines = False
    for i, w in enumerate([7, 8.43, 13, 14, 15.77, 37.88, 18.66, 5.32], 1):
        ws.column_dimensions[X.col_letter(i)].width = w
    merges = []

    def mg(r1, c1, r2, c2):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        merges.append((r1, c1, r2, c2))
    n = len(rows)
    data_start, data_end = 7, 6 + n
    mg(1, 1, 1, 8)
    ws.row_dimensions[1].height = 55.5
    ws.cell(1, 1, title).font = X.font(22, True)
    ws.cell(1, 1).font = X.Font(name="굴림체", size=22, bold=True)
    ws.cell(1, 1).alignment = X.align()
    mg(2, 1, 4, 5)
    mg(2, 6, 4, 6)
    mg(2, 7, 3, 7)
    mg(2, 8, 4, 8)
    for r, h in ((2, 13.5), (3, 13.5), (4, 22.5)):
        ws.row_dimensions[r].height = h
    cyan = X.fill("CCFFFF")

    def box(r, c, val, size=14, h="center", wrap=True, fname=X.FONT):
        cell = ws.cell(r, c, val)
        cell.font = X.Font(name=fname, size=size, bold=True)
        cell.alignment = X.align(h, "center", wrap)
        cell.fill = cyan
    box(2, 1, "지 출 명 세 서")
    box(2, 6, "출장명령부\n대조자")
    box(2, 7, "직급 : " + ap_rank, h="left", wrap=False)
    box(2, 8, "(인)", fname="돋움")
    box(4, 7, "성명 : " + ap_name, h="left", wrap=False)
    mg(5, 7, 5, 8)
    ws.row_dimensions[5].height = 30.75
    for i, h in enumerate(["연번", "성명", "금액(원)", "은행명", "계좌번호", "출장기간(일)", "출장목적"]):
        box(5, i + 1, h, wrap=(i != 0))
    ws.row_dimensions[6].height = 39
    ws.cell(6, 2, "계").font = X.font(11, True)
    ws.cell(6, 2).alignment = X.align()
    c6 = ws.cell(6, 3, grand)
    c6.font = X.Font(name="굴림체", size=11, bold=True)
    c6.alignment = X.align()
    c6.number_format = NUM_FMT
    for idx, (name, amt, bank, acc, days, purpose) in enumerate(rows):
        r = data_start + idx
        ws.row_dimensions[r].height = 39
        ws.cell(r, 1, idx + 1).font = X.Font(name="돋움", size=11, bold=True)
        ws.cell(r, 1).alignment = X.align()
        ws.cell(r, 2, name).font = X.font(11)
        ws.cell(r, 2).alignment = X.align()
        c = ws.cell(r, 3, amt)
        c.font = X.Font(name="돋움", size=11)
        c.alignment = X.align("right")
        c.number_format = NUM_FMT
        ws.cell(r, 4, bank).font = X.font(11)
        ws.cell(r, 4).alignment = X.align()
        ws.cell(r, 5, acc).font = X.Font(name="돋움", size=10)
        ws.cell(r, 5).alignment = X.align()
        ws.cell(r, 6, days).font = X.Font(name="돋움", size=11)
        ws.cell(r, 6).alignment = X.align(wrap=True)
        mg(r, 7, r, 8)
        ws.cell(r, 7, purpose).font = X.font(11)
        ws.cell(r, 7).alignment = X.align(wrap=True)
    longest = max(len(f"{v:,}") for v in [grand] + [r[1] for r in rows])
    if longest + 4 > ws.column_dimensions["C"].width:
        ws.column_dimensions["C"].width = longest + 4
    for r in range(1, data_end + 1):
        for c in range(1, 9):
            top = bottom = left = right = thin
            if r == 1:
                top = medium
            if r == data_end:
                bottom = medium
            if c == 1:
                left = medium
            if c == 8:
                right = medium
            if r >= 2:
                for (r1, c1, r2, c2) in merges:
                    if r1 <= r <= r2 and c1 <= c <= c2:
                        if c > c1:
                            left = Side()
                        if c < c2:
                            right = Side()
                        if r > r1:
                            top = Side()
                        if r < r2:
                            bottom = Side()
            ws.cell(r, c).border = Border(top=top, bottom=bottom, left=left, right=right)
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = 9
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return wb


EHOJO_HEADERS = [
    "거래처명\n(100자리이하)\n(필수입력)", "거래처구분\n(필수입력)", "사업자등록번호\n(10자리)", "생년월일\n(8자리)",
    "입금유형\n(선택)", "은행\n(선택)", "예금주명\n(100자리이하)", "계좌번호\n(30자리이하)",
    "공급가액\n(17자리이하)\n(필수입력)", "부가세액\n(17자리이하)", "공제액\n(금액 0원으로 고정)", "지출액\n(17자리이하)",
    "지로번호\n(19자리이하)", "CMS번호\n(20자리이하)", "대표채주\n(Y/N)", "대표자명\n(100자리)", "우편번호\n(5자리)",
    "기본주소\n(75자이하)", "상세주소\n(75자이하)", "전화번호\n(30자리이하)", "휴대폰번호\n(20자리이하)", "입금명세(100자리)",
]
EHOJO_WIDTHS = [22.625, 21.25, 22.75, 21.625, 16.25, 22.625, 20.625, 26.5, 17.75, 17.75, 23.75, 19.25, 21.625, 8.43, 15.25, 16.125,
                15.75, 25.625, 20.625, 20.625, 8.43, 25.625]


def build_ehojo(rows: list, template_path: str | None = None) -> tuple:
    """rows: [(성명, 은행, 예금주, 계좌, 금액)].  반환: (Workbook, 원본 양식을 썼는지)"""
    wb = ws = None
    using = False
    if template_path:
        try:
            wb = load_workbook(template_path)
            ws = next((s for s in wb.worksheets if "거래처" in s.title), wb.worksheets[0])
            for r in range(2, max(ws.max_row, 1) + 1):
                for c in range(1, 23):
                    ws.cell(r, c).value = None
            using = True
        except Exception:
            wb = None
    if wb is None:
        wb = Workbook()
        ws = wb.active
        ws.title = "거래처입력"
        ws.sheet_view.showGridLines = False
        for i, w in enumerate(EHOJO_WIDTHS, 1):
            ws.column_dimensions[X.col_letter(i)].width = w
        ws.row_dimensions[1].height = 45.75
        for i, h in enumerate(EHOJO_HEADERS, 1):
            c = ws.cell(1, i, h)
            c.font = X.Font(name="돋움", size=12)
            c.alignment = X.align(wrap=True)
            c.fill = X.fill("D9D9D9")
            c.border = X.border("thin", "thin", "thin", "thin")
    for idx, (name, bank, holder, acc, amt) in enumerate(rows):
        r = idx + 2
        if not using:
            ws.row_dimensions[r].height = 16.5

        def put(col, val, size=11, h="center", numfmt=None, fname=X.FONT):
            cell = ws.cell(r, col, val)
            if not using:
                cell.font = X.Font(name=fname, size=size)
                cell.alignment = X.align(h)
                cell.border = X.border("thin", "thin", "thin", "thin")
            if numfmt:
                cell.number_format = numfmt
        put(1, name)
        put(2, "30:개인")
        put(5, "10:계좌이체")
        put(6, bank_to_code(bank))
        put(7, holder or name)
        put(8, acc, 10, fname="돋움")
        put(9, amt, h="right", numfmt=NUM_FMT)
        put(11, 0, h="right", numfmt=NUM_FMT)
        if not using:
            for c in range(1, 23):
                if c not in (1, 2, 5, 6, 7, 8, 9, 11):
                    ws.cell(r, c).border = X.border("thin", "thin", "thin", "thin")
    if not using:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        if rows:
            longest = max(len(f"{r[4]:,}") for r in rows)
            if longest + 4 > ws.column_dimensions["I"].width:
                ws.column_dimensions["I"].width = longest + 4
    return wb, using
