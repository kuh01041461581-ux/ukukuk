# -*- coding: utf-8 -*-
"""급식비(매식비) 정산 — 규칙만 담은 순수 계산 코드 (화면과 무관)

 · 대상 판정: 휴일구분 = 평일 이고 퇴근(실제) 20:00 이후 인 (직원 × 날짜)
 · 배분: 급식처에 입력한 급식인원 합계만큼만, 날짜별 가능 인원에 비례(최대잔여법)해 O 표시
 · 결과: 급식확인서 · 급식지급 상세내역 · 초과근무내역 3개 시트 엑셀
"""
import calendar
import datetime as dt
import re
from dataclasses import dataclass, field
from openpyxl import Workbook

from ...core import excel as X

UNIT_PRICE = 9000


@dataclass
class Record:
    no: object
    y: int
    m: int
    d: int
    date_str: str
    dept: object
    rank: object
    name: str
    holiday: str
    checkin: object
    checkout: object
    checkout_min: int | None
    otmin: object
    detail: object


@dataclass
class Overtime:
    records: list
    employees: list              # 초과근무 내역에 처음 등장한 순서
    meal_qualify: dict           # 이름 → {날짜문자열}
    holidays: set                # 휴일구분=휴일 날짜
    year: int
    month: int
    days: int
    max_meal: int
    master: list = field(default_factory=list)   # 시트2 명단(기존양식 → 없는 사람은 뒤에)


class ParseError(Exception):
    pass


def _parse_date(v):
    if isinstance(v, dt.datetime):
        return v.year, v.month, v.day
    if isinstance(v, dt.date):
        return v.year, v.month, v.day
    s = str(v).strip()
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s) or re.match(r"^(\d{4})\.(\d{1,2})\.(\d{1,2})", s)
    return (int(m.group(1)), int(m.group(2)), int(m.group(3))) if m else None


def _parse_minutes(v):
    if isinstance(v, dt.datetime):
        return v.hour * 60 + v.minute
    if isinstance(v, dt.time):
        return v.hour * 60 + v.minute
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        frac = v % 1
        return round(frac * 24 * 60)
    m = re.match(r"^(\d{1,2}):(\d{2})", str(v).strip())
    return int(m.group(1)) * 60 + int(m.group(2)) if m else None


def parse_overtime(sheets: dict) -> Overtime:
    name = next((n for n in sheets if "초과근무" in n), next(iter(sheets)))
    rows = sheets[name]
    hdr_idx = -1
    for i, r in enumerate(rows[:10]):
        if any(str(c).strip() == "근무일자" for c in r):
            hdr_idx = i
            break
    if hdr_idx < 0:
        raise ParseError('"근무일자" 컬럼을 찾을 수 없습니다. 인사랑에서 내려받은 초과근무 내역 엑셀이 맞는지 확인해주세요.')
    header = [str(c).strip() for c in rows[hdr_idx]]

    def find(n):
        return header.index(n) if n in header else -1
    col = dict(no=find("No."), date=find("근무일자"), dept=find("부서"), rank=find("직급"), name=find("성명"),
               holiday=find("휴일구분"), checkin=find("출근(실제)"), checkout=find("퇴근(실제)"),
               otmin=find("총초과(분)"), detail=find("근무내역"))
    missing = [k for k in ("date", "name", "holiday", "checkout") if col[k] < 0]
    if missing:
        raise ParseError("필수 컬럼이 누락되었습니다: " + ", ".join(missing))

    def cell(r, k):
        return r[col[k]] if col[k] >= 0 and col[k] < len(r) else ""
    records = []
    for r in rows[hdr_idx + 1:]:
        if not r:
            continue
        nm, dv = cell(r, "name"), cell(r, "date")
        if nm in ("", None) or dv in ("", None):
            continue
        d = _parse_date(dv)
        if not d:
            continue
        records.append(Record(
            no=cell(r, "no"), y=d[0], m=d[1], d=d[2], date_str=f"{d[0]}-{d[1]:02d}-{d[2]:02d}",
            dept=cell(r, "dept"), rank=cell(r, "rank"), name=str(nm).strip(), holiday=str(cell(r, "holiday")).strip(),
            checkin=cell(r, "checkin"), checkout=cell(r, "checkout"), checkout_min=_parse_minutes(cell(r, "checkout")),
            otmin=cell(r, "otmin"), detail=cell(r, "detail")))
    if not records:
        raise ParseError("유효한 데이터 행을 찾지 못했습니다.")
    freq: dict = {}
    for rec in records:
        freq[(rec.y, rec.m)] = freq.get((rec.y, rec.m), 0) + 1
    (year, month), _ = max(freq.items(), key=lambda kv: kv[1])
    qualify: dict = {}
    for rec in records:
        if rec.holiday == "평일" and rec.checkout_min is not None and rec.checkout_min >= 20 * 60:
            qualify.setdefault(rec.name, set()).add(rec.date_str)
    employees: list = []
    seen = set()
    for rec in records:
        if rec.name not in seen:
            seen.add(rec.name)
            employees.append(rec.name)
    holidays = {rec.date_str for rec in records if rec.holiday == "휴일"}
    ot = Overtime(records, employees, qualify, holidays, year, month, calendar.monthrange(year, month)[1],
                  sum(len(s) for s in qualify.values()))
    ot.master = list(employees)
    return ot


def parse_template_names(sheets: dict) -> list | None:
    """기존 정산 엑셀의 '급식지급 상세내역' 시트에서 직원 명단·순서를 읽는다"""
    name = (next((n for n in sheets if "급식지급" in n), None) or next((n for n in sheets if "상세내역" in n), None)
            or (list(sheets)[1] if len(sheets) > 1 else list(sheets)[0]))
    rows = sheets.get(name)
    if not rows:
        return None
    hdr = -1
    for i, r in enumerate(rows[:15]):
        if any(str(c).strip() == "성명" for c in r):
            hdr = i
            break
    if hdr < 0:
        return None
    header = [str(c).strip() for c in rows[hdr]]
    nc = header.index("성명")
    names, seen = [], set()
    for r in rows[hdr + 1:]:
        if not r or nc >= len(r):
            continue
        try:
            float(str(r[0]))
        except ValueError:
            continue                      # 연번(첫 칸)이 숫자인 행만 직원 행
        nm = str(r[nc]).strip()
        if nm and nm not in seen:
            names.append(nm)
            seen.add(nm)
    return names


def apply_master(ot: Overtime, template_names: list | None):
    """명단 = 기존양식 명단 우선, 초과근무에만 있는 직원은 뒤에"""
    if template_names:
        master = list(template_names)
        for n in ot.employees:
            if n not in master:
                master.append(n)
    else:
        master = list(ot.employees)
    ot.master = master


def allocate(ot: Overtime, target: int) -> dict:
    """급식인원 target 만큼만 날짜별로 나눠 O 표시. 반환: 이름 → {날짜문자열}"""
    date_map: dict = {}
    for name, dates in ot.meal_qualify.items():
        for d in dates:
            date_map.setdefault(d, []).append(name)
    dates = sorted(date_map)
    cap_total = sum(len(date_map[d]) for d in dates)
    goal = max(0, min(target or 0, cap_total))
    plan = [{"date": d, "cap": len(date_map[d]), "alloc": 0, "frac": 0.0} for d in dates]
    if goal > 0 and cap_total > 0:
        for p in plan:
            exact = goal * p["cap"] / cap_total
            p["alloc"] = int(exact)
            p["frac"] = exact - p["alloc"]
        remainder = goal - sum(p["alloc"] for p in plan)
        by_frac = sorted(plan, key=lambda p: -p["frac"])
        guard = 0
        while remainder > 0 and guard < len(by_frac) * 4 + 20:
            p = by_frac[guard % len(by_frac)]
            if p["alloc"] < p["cap"]:
                p["alloc"] += 1
                remainder -= 1
            guard += 1
    order = ot.master or ot.employees
    rank = {n: i for i, n in enumerate(order)}
    usage: dict = {}
    selected: dict = {}
    for p in plan:
        if p["alloc"] <= 0:
            continue
        cands = sorted(date_map[p["date"]], key=lambda n: (usage.get(n, 0), rank.get(n, 999999)))
        for n in cands[:p["alloc"]]:
            selected.setdefault(n, set()).add(p["date"])
            usage[n] = usage.get(n, 0) + 1
    return selected


def is_red_day(ot: Overtime, day: int) -> bool:
    wd = dt.date(ot.year, ot.month, day).weekday()      # 월=0 … 일=6
    ds = f"{ot.year}-{ot.month:02d}-{day:02d}"
    return wd == 6 or (wd != 5 and ds in ot.holidays)


def is_blue_day(ot: Overtime, day: int) -> bool:
    return dt.date(ot.year, ot.month, day).weekday() == 5


# ─────────────────────── 엑셀 만들기 ───────────────────────
def build_workbook(ot: Overtime, vendors: list, officer_rank: str = "", officer_name: str = "") -> Workbook:
    """vendors: [(급식처, 인원), …]"""
    year, month, days = ot.year, ot.month, ot.days
    total = sum(c for _, c in vendors)
    qualify = allocate(ot, total)
    employees = ot.master or ot.employees
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "급식확인서"
    ws1.sheet_view.showGridLines = False
    for i, w in enumerate([5.11, 20.34, 20.52, 9.99, 5.99, 14.87, 15.99, 0.76], 1):
        ws1.column_dimensions[X.col_letter(i)].width = w
    ws1.row_dimensions[1].height = 4.5
    ws1.merge_cells("A2:G2")
    ws1["A2"] = f"{year}년 {month}월 급식확인서"
    ws1["A2"].font = X.font(20, True)
    ws1["A2"].alignment = X.align()
    ws1.row_dimensions[2].height = 28.5
    ws1["G3"] = "(단위 : 명/원)"
    ws1["G3"].font = X.font(12)
    ws1["G3"].alignment = X.align("right")
    ws1.row_dimensions[3].height = 27
    for i, h in enumerate(["연번", "기간", "급식처", "급식인원", "단가", "지급액", "비고"], 1):
        c = ws1.cell(4, i, h)
        c.font = X.font(12, True)
        c.alignment = X.align()
        c.fill = X.fill("FFFF99")
        c.border = X.border("medium", "double", "medium" if i == 1 else "thin", "medium" if i == 7 else "thin")
    ws1.row_dimensions[4].height = 35.1
    n = len(vendors)
    last = 5 + n
    totals = ["", "계", f"=COUNTA(C6:C{last})", f"=SUM(D6:D{last})", "", f"=SUM(F6:F{last})", ""]
    for i, v in enumerate(totals, 1):
        c = ws1.cell(5, i, v)
        c.font = X.font(12, True)
        c.alignment = X.align("right" if i == 6 else "center")
        c.fill = X.fill("CCFFFF")
        c.border = X.border("double", "medium", "medium" if i == 1 else "thin", "medium" if i == 7 else "thin")
    ws1.row_dimensions[5].height = 35.1
    for idx, (vname, vcount) in enumerate(vendors):
        r = 6 + idx
        vals = [idx + 1, f"{year}.{month}.", vname, vcount, UNIT_PRICE, f"=D{r}*{UNIT_PRICE}", ""]
        for i, v in enumerate(vals, 1):
            c = ws1.cell(r, i, v)
            c.font = X.font(12)
            c.alignment = X.align("right" if i == 6 else "center")
            c.border = X.border("thin", "medium" if idx == n - 1 else "thin", "medium" if i == 1 else "thin", "medium" if i == 7 else "thin")
        ws1.row_dimensions[r].height = 35.1
    note = 6 + n + 1
    ws1.merge_cells(f"A{note}:G{note}")
    ws1[f"A{note}"] = "상기와 같이 식사를 제공하였음을 확인함."
    ws1[f"A{note}"].font = X.font(14)
    ws1[f"A{note}"].alignment = X.align("right")
    ws1.row_dimensions[note].height = 24.75
    dr = note + 2
    for r, text, h, al, ht in ((dr, f"{year}.  {month}.     .", 14, "center", 24.75), (dr + 1, "확인자 : 일상경비출납원", 14, "right", 29.25),
                               (dr + 2, f" {officer_rank} {officer_name} (인)", 14, "right", 31.5)):
        ws1.merge_cells(f"B{r}:G{r}")
        ws1[f"B{r}"] = text
        ws1[f"B{r}"].font = X.font(h)
        ws1[f"B{r}"].alignment = X.align(al)
        ws1.row_dimensions[r].height = ht

    # ── 시트2 ──
    ws2 = wb.create_sheet("급식지급 상세내역")
    ws2.sheet_view.showGridLines = False
    first, lastd = 4, 4 + days - 1
    totc, chkc = lastd + 1, lastd + 2
    for c, w in ((1, 3.76), (2, 8.5), (3, 5.2)):
        ws2.column_dimensions[X.col_letter(c)].width = w
    for c in range(first, lastd + 1):
        ws2.column_dimensions[X.col_letter(c)].width = 3.3
    ws2.column_dimensions[X.col_letter(totc)].width = 5.2
    ws2.column_dimensions[X.col_letter(chkc)].width = 9
    ws2.merge_cells(start_row=1, start_column=1, end_row=2, end_column=lastd)
    ws2.cell(1, 1, f"{year}년 {month}월 급식지급 상세내역").font = X.font(16, True)
    ws2.cell(1, 1).alignment = X.align()
    ws2.row_dimensions[1].height = 26.25
    ws2.row_dimensions[2].height = 14.25
    ws2.row_dimensions[3].height = 39
    ws2.cell(3, 1, "연번")
    ws2.cell(3, 2, "성명")
    ws2.cell(3, 3, "일자\n계")
    for day in range(1, days + 1):
        ws2.cell(3, first + day - 1, day)
    for c in range(1, lastd + 1):
        cell = ws2.cell(3, c)
        color = None
        if c >= first:
            day = c - first + 1
            color = "FF0000" if is_red_day(ot, day) else ("0000FF" if is_blue_day(ot, day) else None)
        cell.font = X.font(9, True, color)
        cell.fill = X.fill("CCFFFF")
        cell.alignment = X.align(wrap=(c == 3))
        cell.border = X.border("medium", "double", "medium" if c == 1 else "thin", "thin")
    emp_last = 4 + len(employees)
    ws2.row_dimensions[4].height = 27.75
    ws2.cell(4, 2, f'=COUNTA(B5:B{emp_last})&"명"')
    ws2.cell(4, 3, f"=SUM(C5:C{emp_last})")
    for day in range(1, days + 1):
        L = X.col_letter(first + day - 1)
        ws2.cell(4, first + day - 1, f'=COUNTIF({L}5:{L}{emp_last},"O")')
    for c in range(1, lastd + 1):
        cell = ws2.cell(4, c)
        cell.fill = X.fill("FFFF99")
        cell.font = X.font(9, True)
        cell.alignment = X.align()
        cell.border = X.border("double", "medium", "medium" if c == 1 else "thin", "thin")
    ld = X.col_letter(lastd)
    ws2.cell(4, totc, f"=SUM(D4:{ld}4)")
    ws2.cell(4, chkc, f'=IF({X.col_letter(totc)}4=C4,"Ok","안맞아요")')
    ws2.cell(4, chkc).font = X.font(9, True, "0000FF")
    ws2.cell(4, chkc).alignment = X.align()
    for idx, emp in enumerate(employees):
        rr = 5 + idx
        ws2.row_dimensions[rr].height = 20
        a = ws2.cell(rr, 1, idx + 1)
        a.alignment = X.align()
        a.border = X.border("thin", "thin", "medium", "thin")
        b = ws2.cell(rr, 2, emp)
        b.alignment = X.align()
        b.border = X.border("thin", "thin", "thin", "thin")
        cc = ws2.cell(rr, 3, f'=COUNTIF(D{rr}:{ld}{rr},"O")')
        cc.fill = X.fill("CCFFFF")
        cc.alignment = X.align()
        cc.border = X.border("thin", "thin", "thin", "thin")
        qs = qualify.get(emp, set())
        for day in range(1, days + 1):
            c = first + day - 1
            ds = f"{year}-{month:02d}-{day:02d}"
            cell = ws2.cell(rr, c, "O" if ds in qs else "")
            cell.font = X.font(9, True, "2F6FED")
            cell.fill = X.fill("FFFFFF")
            cell.alignment = X.align()
            cell.border = X.border("thin", "thin", "thin", "medium" if c == lastd else "thin")

    # ── 시트3 ──
    ws3 = wb.create_sheet("초과근무내역")
    ws3.sheet_view.showGridLines = False
    for i, w in enumerate([6.5, 12, 14, 20, 10, 8, 9, 9, 10, 32], 1):
        ws3.column_dimensions[X.col_letter(i)].width = w
    ws3.merge_cells("A1:J1")
    ws3["A1"] = f"초과근무 내역({year}. {month}.)"
    ws3["A1"].font = X.font(18, True)
    ws3["A1"].alignment = X.align()
    ws3.row_dimensions[1].height = 44.25
    for i, h in enumerate(["No.", "근무일자", "부서", "직급", "성명", "휴일구분", "출근(실제)", "퇴근(실제)", "총초과(분)", "근무내역"], 1):
        c = ws3.cell(2, i, h)
        c.font = X.font(11, True)
        c.alignment = X.align()
        c.fill = X.fill("FFFFCC")
        c.border = X.border("thin", "thin", "thin", "thin")
    ws3.row_dimensions[2].height = 24
    aligns = ["center", "center", "left", "left", "left", "center", "center", "center", "right", "left"]
    for idx, rec in enumerate(sorted(ot.records, key=lambda r: r.date_str)):
        rr = idx + 3
        ws3.row_dimensions[rr].height = 22.5
        vals = [idx + 1, rec.date_str, rec.dept, rec.rank, rec.name, rec.holiday, _cellv(rec.checkin), _cellv(rec.checkout), rec.otmin, rec.detail]
        for i, v in enumerate(vals, 1):
            c = ws3.cell(rr, i, v)
            c.font = X.font(10)
            c.alignment = X.align(aligns[i - 1])
            c.border = X.border(None, "thin", "thin", "thin")
    return wb


def _cellv(v):
    if isinstance(v, dt.time):
        return v.strftime("%H:%M")
    if isinstance(v, dt.datetime):
        return v.strftime("%H:%M")
    return v
