# -*- coding: utf-8 -*-
"""업무(폴더·제출 현황·양식 배포·취합·보고) 시험. 실행: python tests/test_work.py  (openpyxl, olefile 필요)
모든 문서는 지어낸 내용으로 시험 때마다 새로 만든다."""
import datetime as dt
import os
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path[:0] = [os.path.dirname(HERE), HERE]

import openpyxl  # noqa: E402

import core  # noqa: E402
import fixtures as F  # noqa: E402
import gongmun  # noqa: E402
import hwpdoc  # noqa: E402
import work  # noqa: E402

D = dt.date
TODAY = D(2026, 10, 13)
TMP = tempfile.mkdtemp(prefix="dc_work_")
P = lambda *a: os.path.join(TMP, *a)  # noqa: E731

GONGMUN = [
    "도시재생과",
    "수신자 수신자 참조",
    "제목 2027년 도시재생 뉴딜 수요조사 제출 요청",
    "1. 관련: 도시재생과-1234(2026. 10. 1.)",
    "2. 2027년 사업 수요를 붙임 양식에 따라 작성하여 2026. 10. 20.(화)까지 제출하여 주시기 바랍니다.",
    "가. 설명회: 2026. 10. 15.(목) 14:00, 시청 3층 회의실",
    "붙임 수요조사 양식 1부. 끝.",
    "수신자 건축과장, 도로과장, 행정지원과장, 중앙동장",
]


def test_hwp_gongmun():
    F.make_hwp(P("공문.hwp"), GONGMUN)
    text = hwpdoc.read_text(P("공문.hwp"))
    title, items = gongmun.parse(text, TODAY)
    assert title.startswith("2027년 도시재생 뉴딜 수요조사"), title
    dates = {i["date"] for i in items}
    assert D(2026, 10, 20) in dates and D(2026, 10, 15) in dates, items
    assert work.due_from_items(items) == D(2026, 10, 20)
    assert work.recipients(text) == ["건축과", "도로과", "행정지원과", "중앙동"], work.recipients(text)
    assert work.doc_no(text) == "도시재생과-1234"
    F.make_hwpx(P("공문.hwpx"), GONGMUN)
    assert gongmun.parse(hwpdoc.read_text(P("공문.hwpx")), TODAY)[0] == title


def test_match():
    names = ["건축과", "도로과", "행정지원과", "중앙동"]
    assert work.match_target("수요조사_건축과.xlsx", names) == "건축과"
    assert work.match_target("[행정지원] 수요조사 제출.xlsx", names) == "행정지원과"
    assert work.match_target("중앙동 행정복지센터 회신.hwp", names) == "중앙동"
    assert work.match_target("최종본.xlsx", names) is None
    assert work.split_names("건축과, 도로과\n행정지원과·중앙동, 건축과") == names


def make_list_form(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "수요조사"
    ws["A1"] = "2027년 도시재생 수요조사"
    ws.merge_cells("A1:D1")
    ws["A2"] = "제출 부서: {부서}"
    ws.append(["연번", "사업명", "위치", "사업비(백만원)"])         # 3행: 머리글
    for _ in range(3):                                               # 4~6행: 빈 줄 3개
        ws.append([None, None, None, None])
    ws["A7"] = "합계"
    ws.merge_cells("A7:C7")
    ws["D7"] = "=SUM(D4:D6)"
    wb.save(path)


def fill_list(src, dst, rows):
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    for k, r in enumerate(rows):
        for j, v in enumerate(r):
            ws.cell(row=4 + k, column=1 + j).value = v
    wb.save(dst)


def make_grid_form(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "집계"
    ws.append(["구분", "인원", "예산", "비고"])
    ws.append(["청년", None, None, None])
    ws.append(["중장년", None, None, None])
    ws.append(["어르신", None, None, None])
    ws.append(["계", "=SUM(B2:B4)", "=SUM(C2:C4)", None])
    wb.save(path)


def fill_grid(src, dst, vals, note=None):
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    for k, (a, b) in enumerate(vals):
        ws.cell(row=2 + k, column=2).value = a
        ws.cell(row=2 + k, column=3).value = b
    if note:
        ws["D3"] = note
    wb.save(dst)


def test_case_flow():
    data = {"_deleted": {}}
    root = P("업무")
    c = work.add_case(data, work.new_case("2027년 도시재생 수요조사", D(2026, 10, 20)))
    work.set_targets(c, ["건축과", "도로과", "행정지원과", "중앙동"])
    path = work.ensure_folder(root, c)
    assert path.endswith(os.path.join("2026", "1020 2027년 도시재생 수요조사")), path
    for s in work.SUBDIRS:
        assert os.path.isdir(os.path.join(path, s))
    make_list_form(P("양식.xlsx"))
    work.set_form(c, P("양식.xlsx"))
    out = work.distribute(c)
    assert len(out) == 4 and out[0].endswith("양식_건축과.xlsx")
    assert openpyxl.load_workbook(out[0]).active["A2"].value == "제출 부서: 건축과"
    # 제출 파일 3곳 (도로과는 미제출) — 행정지원과는 폴더에 직접 넣은 경우
    fill_list(out[0], P("수요조사_건축과.xlsx"), [[1, "골목 정비", "중앙로", 120], [2, "쉼터 조성", "북문", 80]])
    fill_list(out[3], P("중앙동_회신.xlsx"), [[1, "경로당 리모델링", "중앙동", 50], [2, "벽화", "시장길", 10],
                                         [3, "주차장", "1길", 300]])
    got = work.receive(c, [P("수요조사_건축과.xlsx"), P("중앙동_회신.xlsx")])
    assert [w for _, w in got] == ["건축과", "중앙동"]
    fill_list(out[2], os.path.join(work.sub(c, 2), "행정지원과 제출.xlsx"), [[1, "청사 리모델링", "본청", 900]])
    assert work.scan(c) == ["행정지원과"]
    assert work.progress(c) == (3, 4) and work.missing(c) == ["도로과"]
    msg = work.reminder_text(c, TODAY)
    assert "10.20.(화)까지" in msg and "D-7" in msg and "도로과" in msg and "건축과" not in msg, msg
    res_path, res = work.merge(c, dt.datetime(2026, 10, 21, 9, 0))
    assert res["rows"] == 6 and not res["bad"], res
    ws = openpyxl.load_workbook(res_path)["수요조사"]
    names = [ws.cell(row=r, column=2).value for r in range(4, 10)]
    assert names == ["골목 정비", "쉼터 조성", "청사 리모델링", "경로당 리모델링", "벽화", "주차장"], names
    assert ws.cell(row=3, column=5).value == "제출처" and ws.cell(row=6, column=5).value == "행정지원과"
    assert ws["A10"].value == "합계" and ws["D10"].value == "=SUM(D4:D9)", (ws["A10"].value, ws["D10"].value)
    assert "A10:C10" in [str(r) for r in ws.merged_cells.ranges]
    st = openpyxl.load_workbook(res_path)["제출 현황"]
    assert st.max_row == 4
    # 동기화로 넘어가는 모양
    exp = core.export_state(data)
    assert exp["work"][c["id"]]["targets"][1]["st"] == 0


def test_grid_merge():
    make_grid_form(P("집계양식.xlsx"))
    fill_grid(P("집계양식.xlsx"), P("g1.xlsx"), [(3, 100), (2, 50), (None, None)])
    fill_grid(P("집계양식.xlsx"), P("g2.xlsx"), [(1, 10), (4, 40), (5, 500)], note="경로당 포함")
    res = work.merge_excel(P("집계양식.xlsx"), [("건축과", P("g1.xlsx")), ("도로과", P("g2.xlsx"))], P("g_out.xlsx"))
    ws = openpyxl.load_workbook(P("g_out.xlsx"))["집계"]
    assert [ws.cell(row=r, column=2).value for r in (2, 3, 4)] == [4, 6, 5]
    assert [ws.cell(row=r, column=3).value for r in (2, 3, 4)] == [110, 90, 500]
    assert ws["B5"].value == "=SUM(B2:B4)" and res["rows"] == 0 and res["texts"] == 1
    notes = openpyxl.load_workbook(P("g_out.xlsx"))["글 모음"]
    assert notes["B2"].value == "도로과" and notes["D2"].value == "중장년" and notes["E2"].value == "경로당 포함"


def test_merge_without_form():
    make_list_form(P("양식2.xlsx"))
    fill_list(P("양식2.xlsx"), P("n1.xlsx"), [[1, "가", "x", 1]])
    fill_list(P("양식2.xlsx"), P("n2.xlsx"), [[1, "나", "y", 2], [2, "다", "z", 3]])
    res = work.merge_excel(None, [("A과", P("n1.xlsx")), ("B과", P("n2.xlsx"))], P("n_out.xlsx"))
    ws = openpyxl.load_workbook(P("n_out.xlsx")).active
    assert res["rows"] == 3 and [ws.cell(row=r, column=2).value for r in (4, 5, 6)] == ["가", "나", "다"]
    assert ws["E4"].value == "A과" and ws["E6"].value == "B과"


def test_merge_texts():
    F.make_hwp(P("t_건축과.hwp"), ["건축과 의견", "의견 없음"])
    F.make_hwpx(P("t_도로과.hwpx"), ["도로과 의견", "보도 정비 필요"])
    F.make_docx(P("t_중앙동.docx"), ["중앙동", "경로당 수리"])
    res = work.merge_texts([("건축과", P("t_건축과.hwp")), ("도로과", P("t_도로과.hwpx")),
                            ("중앙동", P("t_중앙동.docx"))], P("t_out.xlsx"))
    ws = openpyxl.load_workbook(P("t_out.xlsx")).active
    assert res["rows"] == 3 and "보도 정비 필요" in ws["C3"].value
    assert "■ 중앙동" in open(P("t_out.txt"), encoding="utf-8-sig").read()


def test_report_and_templates():
    data = {"_deleted": {}}
    mon = D(2026, 10, 12)
    a = core.new_item("도시재생 설명회 참석", time="14:00", done=True)
    b = core.new_item("예산 자료 검토")
    c = core.new_item("감사 수감", end="2026-10-23")
    r = core.new_item("월간 업무계획 제출", repeat={"f": "m", "n": 1})
    core.add_item(data, "2026-10-14", a)
    core.add_item(data, "2026-10-15", b)
    core.add_item(data, "2026-10-21", c)
    core.add_item(data, "2026-09-16", r)
    case = work.add_case(data, work.new_case("수요조사 취합", D(2026, 10, 20)))
    work.set_targets(case, ["건축과", "도로과"])
    work.set_status(case, "건축과", 1)
    txt = work.weekly(data, TODAY)
    assert "■ 주간 업무 보고 (10.12. ~ 10.16.)" in txt, txt
    assert "○ 도시재생 설명회 참석 (10.14.(수) 14:00)" in txt
    assert "□ 미완료·이월\n  ○ 예산 자료 검토 (10.15.(목))" in txt
    assert "○ 감사 수감 (10.21.~10.23.(금))" in txt and "○ 월간 업무계획 제출 (10.16.(금))" in txt
    assert "수요조사 취합 (10.20.(화) 기한, D-7) — 1/2 제출, 미제출: 도로과" in txt
    m = work.monthly(data, TODAY)
    assert m.startswith("■ 10월 업무 보고 (10.1. ~ 10.31.)") and "11.1." in m
    first = {t[1]: work.template_first(t[2], TODAY) for t in work.TEMPLATES}
    assert first["월간 업무계획 제출"] == (D(2026, 10, 25), "m")
    assert first["업무추진비·법인카드 사용내역 정리"] == (D(2026, 10, 31), "ml")
    assert first["사이버 보안진단의 날 (PC 점검)"] == (D(2026, 10, 21), "mw")
    assert first["분기별 추진실적 보고"] == (D(2027, 1, 5), "q")
    assert first["행정사무감사 자료 준비"] == (D(2026, 10, 25), "y")
    assert first["상·하반기 자체평가 자료 작성"] == (D(2026, 12, 20), "h")
    for _, _, rule in work.TEMPLATES:
        d, code = work.template_first(rule, TODAY)
        rep = core.repeat_from_code(code, d)
        assert core.repeat_hit(rep, d, d), rule
    assert work.rule_label(("nth", 3, 2)) == "매월 셋째 수요일"
    # 업무 지우기: 묶인 일정의 연결도 끊고 삭제 기록
    a["case"] = case["id"]
    work.delete_case(data, case["id"])
    assert case["id"] in data["_deleted"] and "case" not in a and not data["_work"]


def test_hwp_distribution_preview():
    F.make_hwp(P("배포.hwp"), ["무시"], distribution=True, preview="제목 배포용 공문\r\n2026. 11. 2.까지 회신")
    t = hwpdoc.read_text(P("배포.hwp"))
    assert t.startswith("제목 배포용 공문") and "\x00" not in t


if __name__ == "__main__":
    n = 0
    for k, fn in list(globals().items()):
        if k.startswith("test_") and callable(fn):
            fn()
            n += 1
            print("ok", k)
    print(f"{n} tests passed ({TMP})")
