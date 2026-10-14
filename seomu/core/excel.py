# -*- coding: utf-8 -*-
"""엑셀 읽기/쓰기 공통 도우미.
   읽기: .xlsx(openpyxl) · .xls(xlrd) → 시트별 2차원 목록. 날짜는 datetime, 시각은 time, 숫자는 float/int.
   쓰기: openpyxl 서식 도우미"""
import datetime as dt
import os
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

EPOCH_ZERO = dt.datetime(1899, 12, 30)


def read_sheets(path: str) -> dict:
    """{시트이름: [[셀값…], …]}  (시트 순서 유지)"""
    ext = os.path.splitext(path)[1].lower()
    out: dict = {}
    if ext == ".xls":
        import xlrd
        book = xlrd.open_workbook(path)
        for sh in book.sheets():
            rows = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    t, v = cell.ctype, cell.value
                    if t == xlrd.XL_CELL_DATE:
                        d = xlrd.xldate_as_datetime(v, book.datemode)
                        row.append(d.time() if v < 1 else d)
                    elif t == xlrd.XL_CELL_EMPTY or t == xlrd.XL_CELL_BLANK:
                        row.append("")
                    elif t == xlrd.XL_CELL_NUMBER and float(v).is_integer():
                        row.append(int(v))
                    else:
                        row.append(v)
                rows.append(row)
            out[sh.name] = rows
        return out
    wb = load_workbook(path, data_only=True, read_only=False)
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            rows.append(["" if v is None else v for v in r])
        out[ws.title] = rows
    return out


def clean(v) -> str:
    return "" if v is None else str(v).strip().replace("\r", "").replace("\n", "")


# ── 쓰기용 서식 ───────────────────────────────────────────────
FONT = "맑은 고딕"


def font(size=11, bold=False, color=None):
    return Font(name=FONT, size=size, bold=bold, color=color)


def side(style):
    return Side(style=style) if style else Side()


def border(top=None, bottom=None, left=None, right=None):
    return Border(top=side(top), bottom=side(bottom), left=side(left), right=side(right))


def fill(rgb: str):
    return PatternFill("solid", fgColor=rgb)


def align(h="center", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def col_letter(n: int) -> str:
    s = ""
    while n > 0:
        n, m = divmod(n - 1, 26)
        s = chr(65 + m) + s
    return s
