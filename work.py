# -*- coding: utf-8 -*-
"""업무(공문 한 건 = 업무 하나) 관리: 업무 폴더, 제출 현황, 양식 배포·취합, 독촉 메시지, 주간 업무 보고, 반복 업무 모음.
화면과 무관한 순수 로직이라 시험(tests/test_work.py)으로 확인한다. 업무 정보는 data["_work"] 에 들어가 폰과 동기화된다.

업무 모양:
  {"id", "title", "due": "YYYY-MM-DD"|None, "received": "YYYY-MM-DD", "doc_no": "", "sender": "",
   "done": False, "note": "", "folder": "PC 의 업무 폴더 경로", "form": "양식 파일 이름",
   "targets": [{"name": "도시재생과", "st": 0 미제출 | 1 제출 | 2 확인, "ts": ms, "file": "받은 파일 이름"}],
   "created": ms, "updated": ms}"""
import datetime as dt
import os
import re
import shutil

import core

SUBDIRS = ("1. 공문", "2. 양식·배포", "3. 제출받은 파일", "4. 취합 결과")
ST_NAMES = ("미제출", "제출", "확인")
PLACEHOLDERS = ("{부서}", "{제출처}", "{기관}", "{부서명}")
WD = "월화수목금토일"


# ── 업무 만들기·찾기 ──
def new_case(title, due=None, **kw):
    c = {"id": core.new_id(), "title": title.strip() or "새 업무", "due": due.isoformat() if isinstance(due, dt.date) else due,
         "received": dt.date.today().isoformat(), "done": False, "targets": [], "created": core.now_ms(),
         "updated": core.now_ms()}
    c.update(kw)
    return c


def cases(data):
    return data.setdefault("_work", {})


def add_case(data, c):
    cases(data)[c["id"]] = c
    return c


def delete_case(data, cid):
    if cases(data).pop(cid, None) is not None:
        data.setdefault("_deleted", {})[cid] = core.now_ms()
    for _, v in list(data.items()):
        if isinstance(v, dict):
            for t in v.get("todos", []) if "todos" in v else []:
                if t.get("case") == cid:
                    t.pop("case", None)
                    core.touch(t)


def touch(c):
    c["updated"] = core.now_ms()
    return c


def sorted_cases(data, include_done=False):
    out = [c for c in (data.get("_work") or {}).values() if include_done or not c.get("done")]
    out.sort(key=lambda c: (bool(c.get("done")), c.get("due") or "9999", c.get("created", 0)))
    return out


def case_items(data, cid):
    """업무에 묶인 일정들 [(날짜키, 항목)]"""
    out = []
    for key, v in data.items():
        if key in ("_deleted", "_work") or not isinstance(v, dict):
            continue
        for t in v.get("todos", []):
            if t.get("case") == cid:
                out.append((key, t))
    return out


def progress(c):
    ts = c.get("targets") or []
    return sum(1 for t in ts if t.get("st", 0) >= 1), len(ts)


def missing(c):
    return [t["name"] for t in c.get("targets") or [] if t.get("st", 0) == 0]


def dday(c, today=None):
    d = core.parse_date(c.get("due"))
    if not d:
        return None
    return (d - (today or dt.date.today())).days


def dday_label(n):
    if n is None:
        return ""
    return "D-DAY" if n == 0 else (f"D-{n}" if n > 0 else f"D+{-n}")


def fmt_date(d, wd=True):
    """공문식 날짜 10.20.(화)"""
    if isinstance(d, str):
        d = core.parse_date(d)
    if not d:
        return ""
    return f"{d.month}.{d.day}." + (f"({WD[d.weekday()]})" if wd else "")


# ── 대상(제출처) ──
def split_names(text):
    """'도시재생과, 건축과 / 각 동' 같은 글을 이름 목록으로 (줄바꿈·쉼표·가운뎃점 구분)"""
    parts = re.split(r"[\n,，、/;·ㆍ]+", text or "")
    out = []
    for p in parts:
        p = re.sub(r"\s+", " ", p).strip(" -•○")
        if p and p not in out:
            out.append(p)
    return out


def set_targets(c, names):
    old = {t["name"]: t for t in c.get("targets") or []}
    c["targets"] = [old.get(n) or {"name": n, "st": 0} for n in names]
    touch(c)


def set_status(c, name, st, file=None):
    for t in c.get("targets") or []:
        if t["name"] == name:
            t["st"] = st
            t["ts"] = core.now_ms()
            if file is not None:
                t["file"] = file
            elif st == 0:
                t.pop("file", None)
    touch(c)


def recipients(text):
    """공문 끝의 '수신자 도시재생과장, 건축과장, …' 에서 대상 후보를 뽑는다 ('수신자 참조' 줄은 건너뜀)"""
    lines = (text or "").replace("\r", "").split("\n")
    body = None
    for k, l in enumerate(lines):
        m = re.match(r"\s*수신자\s*[:：]?\s*(.*)$", l)
        if not m or "참조" in m.group(1) or not m.group(1).strip():
            continue
        body = [m.group(1)]
        for l2 in lines[k + 1:]:                       # 다음 줄로 이어지는 목록
            if not l2.strip() or re.match(r"\s*(시행|접수|우\s?\d|전화|주소|기안|협조|공개|https?:|\d+\.|[가-하]\.)", l2):
                break
            body.append(l2)
    if not body:
        return []
    names = []
    for n in split_names(",".join(body)):
        n = re.sub(r"\(.*?\)", "", n).strip()
        n = re.sub(r"(?<=[과실국소단원팀센터청군구시동읍면])장$", "", n)   # 도시재생과장 → 도시재생과
        if 2 <= len(n) <= 20 and n not in names:
            names.append(n)
    return names if 1 <= len(names) <= 80 else []


def _norm(s):
    return re.sub(r"[\s_\-().\[\]【】]+", "", s or "").lower()


def match_target(filename, names):
    """파일 이름에 들어 있는 제출처를 찾는다. '도시재생과' 와 '도시재생' 처럼 끝 글자(과·팀·동)를 뺀 이름도 본다"""
    stem = _norm(os.path.splitext(os.path.basename(filename))[0])
    best, blen = None, 0
    for n in names:
        keys = {_norm(n)}
        short = re.sub(r"(과|팀|실|국|동|소|단|센터|사업소)$", "", _norm(n))
        if len(short) >= 2:
            keys.add(short)
        for k in keys:
            if k and k in stem and len(k) > blen:
                best, blen = n, len(k)
    return best


# ── 업무 폴더 ──
def safe_name(s, limit=60):
    s = re.sub(r'[\\/:*?"<>|\r\n\t]+', " ", s or "").strip(" .")
    s = re.sub(r"\s+", " ", s)
    return (s[:limit].rstrip() or "업무")


def default_folder(root, c):
    d = core.parse_date(c.get("due")) or core.parse_date(c.get("received")) or dt.date.today()
    return os.path.join(root, str(d.year), f"{d:%m%d} {safe_name(c['title'])}")


def ensure_folder(root, c):
    path = c.get("folder")
    if not path or not os.path.isdir(path):
        path = default_folder(root, c)
        c["folder"] = path
        touch(c)
    for s in SUBDIRS:
        os.makedirs(os.path.join(path, s), exist_ok=True)
    return path


def sub(c, k):
    return os.path.join(c["folder"], SUBDIRS[k])


def _unique(path):
    if not os.path.exists(path):
        return path
    stem, ext = os.path.splitext(path)
    k = 2
    while os.path.exists(f"{stem} ({k}){ext}"):
        k += 1
    return f"{stem} ({k}){ext}"


def copy_in(src, folder, name=None):
    dst = os.path.join(folder, name or os.path.basename(src))
    if os.path.abspath(src) == os.path.abspath(dst):
        return dst
    dst = _unique(dst)
    shutil.copy2(src, dst)
    return dst


def add_doc(c, path):
    return copy_in(path, sub(c, 0))


def set_form(c, path):
    dst = copy_in(path, sub(c, 1))
    c["form"] = os.path.basename(dst)
    touch(c)
    return dst


def form_path(c):
    if not c.get("form") or not c.get("folder"):
        return None
    p = os.path.join(sub(c, 1), c["form"])
    return p if os.path.exists(p) else None


def _fill_xlsx(src, dst, name):
    """엑셀 양식의 {부서} 자리표시를 채운다. 자리표시가 없으면 원본 그대로 복사(서식 100% 유지)"""
    import openpyxl
    wb = openpyxl.load_workbook(src)
    hits = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and any(p in v for p in PLACEHOLDERS):
                    for p in PLACEHOLDERS:
                        v = v.replace(p, name)
                    cell.value = v
                    hits += 1
    if hits:
        wb.save(dst)
    else:
        shutil.copy2(src, dst)
    return hits


def distribute(c):
    """제출처마다 '양식이름_제출처.확장자' 파일을 만든다 (자리표시 {부서} 가 있으면 이름을 채워 넣음)"""
    import hwpdoc
    src = form_path(c)
    if not src:
        raise ValueError("먼저 양식 파일을 넣어 주세요")
    out_dir = os.path.join(sub(c, 1), "배포용")
    os.makedirs(out_dir, exist_ok=True)
    stem, ext = os.path.splitext(c["form"])
    out = []
    for t in c.get("targets") or []:
        dst = os.path.join(out_dir, f"{safe_name(stem, 40)}_{safe_name(t['name'], 30)}{ext}")
        if ext.lower() in (".xlsx", ".xlsm"):
            _fill_xlsx(src, dst, t["name"])
        elif ext.lower() == ".hwpx":
            hwpdoc.fill_hwpx(src, dst, {k.strip("{}"): t["name"] for k in PLACEHOLDERS})
        else:
            shutil.copy2(src, dst)
        out.append(dst)
    return out


def receive(c, paths):
    """받은 제출 파일을 '3. 제출받은 파일' 로 복사하고, 파일 이름으로 제출처를 찾아 '제출' 표시. [(경로, 제출처|None)]"""
    names = [t["name"] for t in c.get("targets") or []]
    out = []
    for p in paths:
        who = match_target(p, names)
        dst = copy_in(p, sub(c, 2))
        if who:
            set_status(c, who, max(1, _st(c, who)), os.path.basename(dst))
        out.append((dst, who))
    return out


def _st(c, name):
    for t in c.get("targets") or []:
        if t["name"] == name:
            return t.get("st", 0)
    return 0


def scan(c):
    """'3. 제출받은 파일' 폴더에 직접 넣은 파일도 찾아서 제출 표시. 새로 표시한 제출처 목록"""
    if not c.get("folder") or not os.path.isdir(sub(c, 2)):
        return []
    names = [t["name"] for t in c.get("targets") or []]
    new = []
    for f in sorted(os.listdir(sub(c, 2))):
        if f.startswith("~$"):
            continue
        who = match_target(f, names)
        if who and _st(c, who) == 0:
            set_status(c, who, 1, f)
            new.append(who)
    return new


def submissions(c):
    """[(제출처, 파일 경로)] — 이름으로 짝지은 파일 + 짝을 못 찾은 파일은 (파일 이름, 경로)"""
    if not c.get("folder") or not os.path.isdir(sub(c, 2)):
        return []
    names = [t["name"] for t in c.get("targets") or []]
    out = []
    for f in sorted(os.listdir(sub(c, 2))):
        p = os.path.join(sub(c, 2), f)
        if f.startswith("~$") or not os.path.isfile(p):
            continue
        out.append((match_target(f, names) or os.path.splitext(f)[0], p))
    order = {n: i for i, n in enumerate(names)}
    out.sort(key=lambda x: order.get(x[0], 9999))
    return out


# ── 취합 ──
def _empty(v):
    return v is None or (isinstance(v, str) and (not v.strip() or any(p in v for p in PLACEHOLDERS)))


def _rows(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


def _shift_formula(f, ins, k):
    """행을 끼워 넣은 뒤 아래쪽 수식의 셀 주소를 고친다. SUM(C5:C14) 처럼 데이터 끝(ins-1)에서 끝나는 범위는 늘린다"""
    def rng(m):
        a, r1, b, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
        if r1 >= ins:
            r1 += k
        if r2 >= ins - 1:
            r2 += k
        return f"{a}{r1}:{b}{r2}"
    f = re.sub(r"(\$?[A-Z]{1,3}\$?)(\d+):(\$?[A-Z]{1,3}\$?)(\d+)", rng, f)
    return re.sub(r"(?<![:A-Z$])(\$?[A-Z]{1,3}\$?)(\d+)(?![:\d(])",
                  lambda m: f"{m.group(1)}{int(m.group(2)) + k if int(m.group(2)) >= ins else m.group(2)}", f)


def _merged(cell):
    return type(cell).__name__ == "MergedCell"


def _has_ph(v):
    return isinstance(v, str) and any(p in v for p in PLACEHOLDERS)


def _blank_t(v):
    """양식 칸이 '빈칸'인가 (자리표시·0·수식은 빈칸으로 본다 — 수식은 따로 지킨다)"""
    return _empty(v) or v == 0 or (isinstance(v, str) and v.startswith("="))


def _common_template(books, name):
    """양식 파일이 없을 때: 모든 제출 파일에서 같은 칸에 같은 값이 있으면 양식(머리글·항목 이름)으로 본다"""
    sheets = [wb[name] for _, wb in books if name in wb.sheetnames]
    if len(sheets) < 2:
        return None
    grids = [_rows(ws) for ws in sheets]
    h = max(len(g) for g in grids)
    w = max((len(r) for g in grids for r in g), default=0)
    T = []
    for i in range(h):
        row = []
        for j in range(w):
            vals = [g[i][j] if i < len(g) and j < len(g[i]) else None for g in grids]
            same = isinstance(vals[0], str) and vals[0].strip() and all(v == vals[0] for v in vals)
            row.append(vals[0] if same else None)               # 숫자는 우연히 같을 수 있어 양식으로 보지 않는다
        T.append(row)
    return T


def merge_excel(form, subs, out_path):
    """엑셀 취합. form: 양식(없으면 제출 파일끼리 비교해 공통 부분을 양식으로), subs: [(제출처, 경로)]
    - 목록형: 양식에서 비어 있던 줄에 적힌 내용 → 제출처 열을 붙여 한 표로 이어 붙인다
    - 집계형: 양식에 항목 이름이 있는 줄의 빈칸에 적힌 숫자 → 같은 칸끼리 더한다 (글은 '글 모음' 시트로)
    결과: {"rows": 이어 붙인 줄 수, "sums": 더한 칸 수, "texts": 글 모음 개수, "bad": [(제출처, 이유)]}"""
    import openpyxl
    from openpyxl.utils import get_column_letter
    res = {"rows": 0, "sums": 0, "texts": 0, "bad": []}
    loaded = []
    for who, p in subs:
        if os.path.splitext(p)[1].lower() not in (".xlsx", ".xlsm"):
            res["bad"].append((who, "엑셀(.xlsx) 파일이 아님"))
            continue
        try:
            loaded.append((who, openpyxl.load_workbook(p, data_only=True)))
        except Exception as e:  # noqa: BLE001
            res["bad"].append((who, f"열 수 없음 ({e.__class__.__name__})"))
    if not form and not loaded:
        raise ValueError("취합할 엑셀 파일이 없어요")
    base = form or [p for who, p in subs if who in {w for w, _ in loaded}][0]
    out = openpyxl.load_workbook(base)
    tmpl = openpyxl.load_workbook(base) if form else None
    notes = []
    for ws_o in out.worksheets:
        name = ws_o.title
        if form:
            T = _rows(tmpl[name])
        else:
            T = _common_template(loaded, name)
            if T is None:
                continue
            for i, row in enumerate(ws_o.iter_rows()):      # 첫 제출 파일의 내용 칸은 비우고 시작
                for j, cell in enumerate(row):
                    tv = T[i][j] if i < len(T) and j < len(T[i]) else None
                    if tv is None and not _merged(cell) and not (isinstance(cell.value, str) and cell.value.startswith("=")):
                        cell.value = None
        width = max([len(r) for r in T] + [1])
        last_col = max([j + 1 for r in T for j, v in enumerate(r) if not _empty(v)] + [1])
        records, sums = [], {}
        for who, wb in loaded:
            ws = wb[name] if name in wb.sheetnames else (wb.worksheets[0] if len(out.worksheets) == 1 else None)
            if ws is None:
                continue
            for i, row in enumerate(_rows(ws)):
                trow = T[i] if i < len(T) else []
                labelled = any(not _blank_t(v) or _has_ph(v) for v in trow)
                vals = [(j, v) for j, v in enumerate(row)       # {부서} 자리표시 칸은 이름만 들어 있으니 뺀다
                        if not _empty(v) and (j >= len(trow) or _blank_t(trow[j]) and not _has_ph(trow[j]))]
                if not vals:
                    continue
                if not labelled:
                    records.append((i, who, row))
                    continue
                for j, v in vals:
                    tv = trow[j] if j < len(trow) else None
                    if isinstance(tv, str) and tv.startswith("="):
                        continue                                   # 양식의 수식 칸은 그대로 두면 다시 계산됨
                    if isinstance(v, (int, float)) and not isinstance(v, bool):
                        sums[(i, j)] = sums.get((i, j), 0) + v
                    else:
                        notes.append((name, who, f"{get_column_letter(j + 1)}{i + 1}", _label_of(T, i), str(v)))
        for (i, j), v in sums.items():
            if not _merged(ws_o.cell(row=i + 1, column=j + 1)):
                ws_o.cell(row=i + 1, column=j + 1).value = v
        res["sums"] += len(sums)
        if not records:
            continue
        # 목록형: 데이터가 시작되는 줄부터 차례로 채운다 (모자라면 합계 줄 앞에 줄을 끼워 넣는다)
        start = min(i for i, _, _ in records)
        ins = _next_filled(T, start)
        room = (ins - start) if ins is not None else 10 ** 6
        need = len(records)
        if need > room:
            k = need - room
            ws_o.insert_rows(ins + 1, k)
            _fix_after_insert(ws_o, ins + 1, k)
        col = max(last_col, max(len(r) for _, _, r in records)) + 1
        if start >= 1:
            h = ws_o.cell(row=start, column=col)
            h.value = "제출처"
            src = ws_o.cell(row=start, column=col - 1)
            if src.has_style:
                h._style = src._style
        tpl = ws_o.cell(row=start + 1, column=col - 1)
        for n, (i, who, row) in enumerate(records):
            r = start + n + 1
            for j, v in enumerate(row):
                cell = ws_o.cell(row=r, column=j + 1)
                if _empty(v) or _merged(cell) or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    continue
                cell.value = v
            c2 = ws_o.cell(row=r, column=col)
            c2.value = who
            if tpl.has_style:
                c2._style = tpl._style
        res["rows"] += need
    st = out.create_sheet("제출 현황")
    st.append(["제출처", "상태", "파일"])
    got = {w for w, _ in loaded}
    for who, p in subs:
        st.append([who, "취합함" if who in got else "제외", os.path.basename(p)])
    for who, why in res["bad"]:
        st.append([who, why, ""])
    st.column_dimensions["A"].width = 18
    st.column_dimensions["C"].width = 40
    if notes:
        ws_n = out.create_sheet("글 모음")
        ws_n.append(["시트", "제출처", "칸", "항목", "적은 내용"])
        for r in notes:
            ws_n.append(list(r))
        res["texts"] = len(notes)
    out.save(out_path)
    return res


def _label_of(T, i):
    if i < len(T):
        for v in T[i]:
            if isinstance(v, str) and v.strip():
                return v.strip()[:30]
    return ""


def _next_filled(T, start):
    for i in range(start + 1, len(T)):
        if any(not _blank_t(v) for v in T[i]):
            return i
    return None


def _fix_after_insert(ws, at, k):
    """openpyxl 은 행을 끼워 넣어도 병합 칸·수식을 옮기지 않으므로 직접 고친다 (at: 1부터 세는 행 번호)"""
    from openpyxl.worksheet.cell_range import CellRange
    moved = []
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= at:
            ws.merged_cells.remove(rng)
            moved.append(CellRange(min_col=rng.min_col, min_row=rng.min_row + k, max_col=rng.max_col,
                                   max_row=rng.max_row + k))
    for r in moved:
        ws.merge_cells(r.coord)
    for row in ws.iter_rows(min_row=at + k):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = _shift_formula(cell.value, at, k)


def merge_texts(subs, out_path):
    """한글·워드·글 파일 취합: 제출처마다 본문 글자를 한 엑셀(제출처 | 내용)과 한 글 파일로 모은다"""
    import hwpdoc
    import openpyxl
    from openpyxl.styles import Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "본문 모음"
    ws.append(["제출처", "파일", "내용"])
    txt, bad = [], []
    for who, p in subs:
        try:
            body = hwpdoc.read_text(p).strip()
        except Exception as e:  # noqa: BLE001
            bad.append((who, str(e)[:60]))
            continue
        ws.append([who, os.path.basename(p), body[:32000]])
        ws.cell(row=ws.max_row, column=3).alignment = Alignment(wrap_text=True, vertical="top")
        txt.append(f"■ {who}\n{body}\n")
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 90
    for who, why in bad:
        ws.append([who, "", "읽지 못함: " + why])
    wb.save(out_path)
    with open(os.path.splitext(out_path)[0] + ".txt", "w", encoding="utf-8-sig") as f:
        f.write("\n".join(txt))
    return {"rows": len(txt), "bad": bad}


def merge(c, now=None):
    """업무의 제출 파일을 취합해 '4. 취합 결과' 에 저장. (결과 경로, 요약)"""
    subs = submissions(c)
    if not subs:
        raise ValueError("'3. 제출받은 파일' 에 받은 파일이 없어요")
    now = now or dt.datetime.now()
    form = form_path(c)
    kind = os.path.splitext(form or subs[0][1])[1].lower()
    out = os.path.join(sub(c, 3), f"취합_{safe_name(c['title'], 40)}_{now:%m%d_%H%M}.xlsx")
    if kind in (".xlsx", ".xlsm"):
        return out, merge_excel(form, [s for s in subs], out)
    if kind == ".xls":
        raise ValueError("예전 엑셀(.xls)은 취합할 수 없어요. 엑셀에서 '다른 이름으로 저장 → .xlsx' 로 바꿔 주세요")
    return out, merge_texts(subs, out)


# ── 메신저로 보낼 글 ──
def reminder_text(c, today=None):
    today = today or dt.date.today()
    miss = missing(c)
    due = core.parse_date(c.get("due"))
    lines = [f"[{c['title']}] 자료 제출 안내"]
    if due:
        n = (due - today).days
        when = "오늘" if n == 0 else ("내일" if n == 1 else fmt_date(due))
        lines.append(f"제출 기한: {fmt_date(due)}까지" + (f" ({when})" if n in (0, 1) else f" ({dday_label(n)})"))
    if miss:
        lines.append("아직 받지 못한 곳: " + ", ".join(miss))
    late = due is not None and due < today
    if not c.get("targets"):
        lines.append("기한이 지났습니다. 확인 후 빠른 제출 부탁드립니다." if late else "기한 내 제출 부탁드립니다.")
    elif miss:
        lines.append("기한이 지났습니다. 확인 후 빠른 제출 부탁드립니다." if late
                     else "기한 내 제출 부탁드립니다. 이미 보내셨다면 양해 부탁드립니다.")
    else:
        lines.append("모든 자료를 받았습니다. 감사합니다.")
    return "\n".join(lines)


def status_text(c, today=None):
    """상급자 보고용 한 줄 요약: '○ 수요조사 취합(10.20.(화) 기한, D-3) — 9/12 제출, 미제출: A, B'"""
    got, tot = progress(c)
    due = core.parse_date(c.get("due"))
    s = f"○ {c['title']}"
    if due:
        s += f" ({fmt_date(due)} 기한, {dday_label(dday(c, today))})"
    if tot:
        s += f" — {got}/{tot} 제출"
        miss = missing(c)
        if miss:
            s += ", 미제출: " + ", ".join(miss[:8]) + (" 외" if len(miss) > 8 else "")
    return s


# ── 주간·월간 업무 보고 ──
def _range_items(data, a, b):
    """a~b 기간에 있는 항목 [(날짜, 항목, 끝났나)] — 반복 일정은 날짜마다"""
    out = []
    d = a
    seen = set()
    while d <= b:
        for _, t, s0 in core.day_items(data, d):
            base = core.base_of(t)
            e0 = core.item_end(base, s0)
            key = (id(base), d if base.get("repeat") else s0)
            if key in seen:
                continue
            seen.add(key)
            out.append((d if base.get("repeat") or not e0 else max(s0, a), t, bool(t.get("done")), s0, e0))
        d += dt.timedelta(days=1)
    return out


def _line(d, t, s0=None, e0=None):
    when = fmt_date(d)
    if e0 and s0:
        when = f"{fmt_date(s0, False)}~{fmt_date(e0)}"
    tm = f" {t['time']}" if t.get("time") else ""
    return f"  ○ {t.get('text', '').strip()} ({when}{tm})"


def report(data, a, b, na, nb, title="주간 업무 보고", today=None):
    """[a,b] 실적 + [na,nb] 계획 + 진행 중인 업무 → 보고용 글"""
    today = today or dt.date.today()
    cur = _range_items(data, a, b)
    nxt = _range_items(data, na, nb)
    done = [x for x in cur if x[2]]
    left = [x for x in cur if not x[2]]
    L = [f"■ {title} ({fmt_date(a, False)} ~ {fmt_date(b, False)})", ""]
    L.append("□ 추진 실적")
    L += [_line(d, t, s0, e0) for d, t, _, s0, e0 in done] or ["  ○ (끝낸 일정 없음)"]
    if left:
        L += ["", "□ 미완료·이월"]
        L += [_line(d, t, s0, e0) for d, t, _, s0, e0 in left]
    L += ["", "□ 향후 계획" + f" ({fmt_date(na, False)} ~ {fmt_date(nb, False)})"]
    L += [_line(d, t, s0, e0) for d, t, _, s0, e0 in nxt] or ["  ○ (예정 일정 없음)"]
    act = sorted_cases(data)
    if act:
        L += ["", "□ 진행 중인 업무"]
        L += ["  " + status_text(c, today) for c in act]
    return "\n".join(L)


def week_bounds(d):
    mon = d - dt.timedelta(days=d.weekday())
    return mon, mon + dt.timedelta(days=4)


def weekly(data, today=None, last=False):
    today = today or dt.date.today()
    a, b = week_bounds(today - dt.timedelta(days=7) if last else today)
    na, nb = a + dt.timedelta(days=7), b + dt.timedelta(days=7)
    return report(data, a, b, na, nb, "주간 업무 보고", today)


def monthly(data, today=None):
    import calendar
    today = today or dt.date.today()
    a = today.replace(day=1)
    b = today.replace(day=calendar.monthrange(today.year, today.month)[1])
    na = b + dt.timedelta(days=1)
    nb = na.replace(day=calendar.monthrange(na.year, na.month)[1])
    return report(data, a, b, na, nb, f"{today.month}월 업무 보고", today)


# ── 반복 업무 모음 (기관마다 날짜가 다르니 넣기 전에 고칠 수 있게 보여 준다) ──
# rule: ("day", 일) 매월 그날 / ("last",) 매월 말일 / ("nth", k, 요일0~6) 매월 k째 요일 /
#       ("q", 일) 분기 첫 달(1·4·7·10월) 그날 / ("h", 월, 일) 반기 / ("y", 월, 일) 매년
TEMPLATES = [
    ("매월", "월간 업무계획 제출", ("day", 25)),
    ("매월", "업무추진비·법인카드 사용내역 정리", ("last",)),
    ("매월", "월례회의(정례조회)", ("nth", 1, 0)),
    ("매월", "사이버 보안진단의 날 (PC 점검)", ("nth", 3, 2)),
    ("매월", "주요 업무 추진실적 입력", ("day", 5)),
    ("분기", "분기별 추진실적 보고", ("q", 5)),
    ("분기", "개인정보 보호 자율점검", ("q", 20)),
    ("반기", "상·하반기 자체평가 자료 작성", ("h", 6, 20)),
    ("매년", "연말정산 서류 제출", ("y", 1, 15)),
    ("매년", "결산 자료 준비", ("y", 2, 10)),
    ("매년", "다음 연도 본예산 요구 자료 준비", ("y", 8, 20)),
    ("매년", "행정사무감사 자료 준비", ("y", 10, 25)),
    ("매년", "업무 인수인계서 정리 (정기 인사 전)", ("y", 1, 2)),
    ("매년", "청렴·폭력예방 등 법정 의무교육 이수 확인", ("y", 11, 30)),
]


def template_first(rule, today=None):
    """반복 업무가 오늘 이후 처음 돌아오는 날과 반복 코드 (date, code)"""
    import calendar
    today = today or dt.date.today()
    kind = rule[0]

    def months(step, first):
        y, m = today.year, today.month
        for _ in range(30):
            if (m - first) % step == 0:
                yield y, m
            m += 1
            if m > 12:
                y, m = y + 1, 1

    if kind == "day":
        for y, m in months(1, 1):
            d = dt.date(y, m, min(rule[1], calendar.monthrange(y, m)[1]))
            if d >= today:
                return d, "m"
    if kind == "last":
        for y, m in months(1, 1):
            d = dt.date(y, m, calendar.monthrange(y, m)[1])
            if d >= today:
                return d, "ml"
    if kind == "nth":
        k, wd = rule[1], rule[2]
        for y, m in months(1, 1):
            first = dt.date(y, m, 1)
            d = first + dt.timedelta(days=(wd - first.weekday()) % 7 + 7 * (k - 1))
            if d.month == m and d >= today:
                return d, "mw"
    if kind in ("q", "h"):
        step = 3 if kind == "q" else 6
        mon = 1 if kind == "q" else rule[1]
        day = rule[1] if kind == "q" else rule[2]
        for y, m in months(step, mon):
            d = dt.date(y, m, min(day, calendar.monthrange(y, m)[1]))
            if d >= today:
                return d, kind
    if kind == "y":
        for y in (today.year, today.year + 1):
            d = dt.date(y, rule[1], min(rule[2], calendar.monthrange(y, rule[1])[1]))
            if d >= today:
                return d, "y"
    return today, ""


def rule_label(rule):
    k = rule[0]
    if k == "day":
        return f"매월 {rule[1]}일"
    if k == "last":
        return "매월 말일"
    if k == "nth":
        return f"매월 {['', '첫째', '둘째', '셋째', '넷째'][rule[1]]} {WD[rule[2]]}요일"
    if k == "q":
        return f"1·4·7·10월 {rule[1]}일"
    if k == "h":
        return f"{rule[1]}·{(rule[1] + 5) % 12 + 1}월 {rule[2]}일"
    if k == "y":
        return f"매년 {rule[1]}월 {rule[2]}일"
    return ""


# ── 공문 → 업무 ──
_DUEWORD = re.compile(r"제출|회신|보고|기한|마감|신청|입력|등록")


def due_from_items(items):
    """공문에서 뽑은 일정 중 '제출·회신' 같은 마감 날짜 (없으면 가장 늦은 날짜)"""
    cand = [i for i in items if _DUEWORD.search(i.get("text", ""))]
    pool = cand or items
    if not pool:
        return None
    return max((i.get("end") or i["date"]) if not i.get("period") else i["date"] for i in pool)


def doc_no(text):
    m = re.search(r"([가-힣A-Za-z]{2,20}-\d{2,7})\s*\(\s*20\d{2}\s*[.\-]\s*\d{1,2}\s*[.\-]\s*\d{1,2}", text or "")
    return m.group(1) if m else ""
